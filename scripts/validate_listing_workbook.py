#!/usr/bin/env python3
"""Validate the fixed nine-column Ozon listing-reference sheet before delivery."""

from __future__ import annotations

import re
import sys
from math import ceil, isfinite
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "sku",
    "自定义货号",
    "重量（g）",
    "长度（mm）",
    "宽度（mm）",
    "高度（mm）",
    "商品详情页链接",
    "成本价",
    "建议售价",
]
DETAIL_URL = re.compile(r"^https://detail\.1688\.com/offer/\d+\.html(?:[?#].*)?$")
REQUIRED_PRICING_HEADERS = [
    "自定义货号",
    "采购价下限(CNY)",
    "采购价上限(CNY)",
    "国内运费下限(CNY)",
    "国内运费上限(CNY)",
    "贴单费(CNY)",
    "其他国内成本(CNY)",
    "成本价下限(CNY)",
    "成本价上限(CNY)",
    "国际运费(CNY)",
    "建议售价(RUB)",
    "定价状态",
    "包装数据来源",
    "包装数据原文/推导",
    "包装数据置信度",
    "缺失字段与补采记录",
]


def has_cny_unit(value: object, number_format: str) -> bool:
    if isinstance(value, (int, float)):
        return "¥" in number_format or "CNY" in number_format.upper()
    text = str(value or "")
    return "¥" in text or "CNY" in text.upper()


def has_rub_unit(number_format: str) -> bool:
    return "₽" in number_format or "RUB" in number_format.upper()


def find_header_row(sheet) -> int:
    for row in range(1, min(sheet.max_row, 50) + 1):
        values = [sheet.cell(row, col).value for col in range(1, 10)]
        if values == EXPECTED_HEADERS:
            return row
    raise ValueError("未找到完全匹配的九列表头")


def find_pricing_header(sheet) -> tuple[int, dict[str, int]]:
    for row in range(1, min(sheet.max_row, 60) + 1):
        mapping = {
            str(sheet.cell(row, col).value).strip(): col
            for col in range(1, sheet.max_column + 1)
            if sheet.cell(row, col).value not in (None, "")
        }
        if all(header in mapping for header in REQUIRED_PRICING_HEADERS):
            return row, mapping
    raise ValueError("定价依据缺少成本拆分必需列")


def as_number(value: object) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    number = float(value)
    return number if isfinite(number) else None


def is_positive_number(value: object) -> bool:
    return (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and isfinite(float(value))
        and float(value) > 0
    )


def completion_summary(rows: list[list[object]]) -> dict[str, object]:
    reported_fields = {
        "重量（g）": 2,
        "长度（mm）": 3,
        "宽度（mm）": 4,
        "高度（mm）": 5,
        "建议售价": 8,
    }
    field_counts = {
        name: sum(is_positive_number(row[index]) for row in rows)
        for name, index in reported_fields.items()
    }
    core_indexes = (reported_fields["重量（g）"], reported_fields["建议售价"])
    complete = sum(
        all(is_positive_number(row[index]) for index in core_indexes)
        for row in rows
    )
    total = len(rows)
    return {
        "total": total,
        "complete": complete,
        "required": ceil(total * 0.90),
        "ratio": (complete / total) if total else 0.0,
        "field_counts": field_counts,
    }


def validate(path: Path) -> list[str]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    if "Ozon上品参考" not in workbook.sheetnames:
        return ["缺少工作表：Ozon上品参考"]
    sheet = workbook["Ozon上品参考"]
    errors: list[str] = []
    try:
        header_row = find_header_row(sheet)
    except ValueError as exc:
        return [str(exc)]

    populated_rows: list[list[object]] = []
    main_rows_by_item: dict[str, tuple[int, list[object]]] = {}
    seen_offer_skus: dict[tuple[str, str], int] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(1, 10)]
        if all(value in (None, "") for value in values):
            continue
        populated_rows.append(values)
        if not values[0] or not values[1]:
            errors.append(f"第{row}行：sku或自定义货号为空")
        else:
            item_no = str(values[1]).strip()
            if item_no in main_rows_by_item:
                errors.append(
                    f"第{row}行：自定义货号{item_no}与第"
                    f"{main_rows_by_item[item_no][0]}行重复"
                )
            else:
                main_rows_by_item[item_no] = (row, values)
        for col in range(3, 7):
            value = values[col - 1]
            if value not in (None, "") and not is_positive_number(value):
                errors.append(
                    f"第{row}行：{sheet.cell(header_row, col).value}必须为空或正数值"
                )
        link = values[6]
        if not isinstance(link, str) or not DETAIL_URL.match(link.strip()):
            errors.append(f"第{row}行：G列不是有效1688详情链接")
        elif values[0]:
            offer_sku = (link.strip(), str(values[0]).strip())
            if offer_sku in seen_offer_skus:
                errors.append(
                    f"第{row}行：同一详情链接与sku组合和第"
                    f"{seen_offer_skus[offer_sku]}行重复"
                )
            else:
                seen_offer_skus[offer_sku] = row
        for col in [3, 4, 5, 6, 8, 9]:
            value = values[col - 1]
            if isinstance(value, str) and value.lower().startswith(("http://", "https://")):
                errors.append(f"第{row}行：链接错误地出现在{sheet.cell(header_row, col).value}列")
        cost_cell = sheet.cell(row, 8)
        if cost_cell.value in (None, ""):
            errors.append(f"第{row}行：成本价为空")
        elif not has_cny_unit(cost_cell.value, cost_cell.number_format or ""):
            errors.append(f"第{row}行：成本价未显示CNY/¥单位")
        price_cell = sheet.cell(row, 9)
        if price_cell.value not in (None, ""):
            if not is_positive_number(price_cell.value):
                errors.append(f"第{row}行：建议售价必须为正数值或留空")
            elif not has_rub_unit(price_cell.number_format or ""):
                errors.append(f"第{row}行：建议售价未显示RUB/₽单位")

    completion = completion_summary(populated_rows)
    if completion["total"] == 0:
        errors.append("Ozon上品参考没有商品SKU数据行")
    elif completion["complete"] < completion["required"]:
        field_text = "，".join(
            f"{name}{count}/{completion['total']}"
            for name, count in completion["field_counts"].items()
        )
        errors.append(
            "重量与建议售价联合完整率不足90%："
            f"完整{completion['complete']}/{completion['total']}"
            f"（{completion['ratio']:.1%}），至少需要{completion['required']}行；"
            f"各字段：{field_text}"
        )

    if "定价依据" not in workbook.sheetnames:
        errors.append("缺少工作表：定价依据")
        return errors
    pricing = workbook["定价依据"]
    try:
        pricing_header_row, columns = find_pricing_header(pricing)
    except ValueError as exc:
        errors.append(str(exc))
        return errors

    pricing_rows_by_item: dict[str, int] = {}
    for row in range(pricing_header_row + 1, pricing.max_row + 1):
        item_no = pricing.cell(row, columns["自定义货号"]).value
        if item_no in (None, ""):
            continue
        item_key = str(item_no).strip()
        if item_key in pricing_rows_by_item:
            errors.append(
                f"定价依据第{row}行：自定义货号{item_key}与第"
                f"{pricing_rows_by_item[item_key]}行重复"
            )
        else:
            pricing_rows_by_item[item_key] = row
        numeric_headers = REQUIRED_PRICING_HEADERS[1:9]
        numbers: dict[str, float] = {}
        for header in numeric_headers:
            value = as_number(pricing.cell(row, columns[header]).value)
            if value is None:
                errors.append(f"定价依据第{row}行：{header}必须为数值")
            else:
                numbers[header] = value
        positive_headers = {
            "采购价下限(CNY)", "采购价上限(CNY)",
            "成本价下限(CNY)", "成本价上限(CNY)",
        }
        nonnegative_headers = {
            "国内运费下限(CNY)", "国内运费上限(CNY)",
            "贴单费(CNY)", "其他国内成本(CNY)",
        }
        for header in positive_headers & numbers.keys():
            if numbers[header] <= 0:
                errors.append(f"定价依据第{row}行：{header}必须大于0")
        for header in nonnegative_headers & numbers.keys():
            if numbers[header] < 0:
                errors.append(f"定价依据第{row}行：{header}不能为负数")
        if (
            {"采购价下限(CNY)", "采购价上限(CNY)"}.issubset(numbers)
            and numbers["采购价下限(CNY)"] > numbers["采购价上限(CNY)"]
        ):
            errors.append(f"定价依据第{row}行：采购价下限高于上限")
        if (
            {"国内运费下限(CNY)", "国内运费上限(CNY)"}.issubset(numbers)
            and numbers["国内运费下限(CNY)"] > numbers["国内运费上限(CNY)"]
        ):
            errors.append(f"定价依据第{row}行：国内运费下限高于上限")
        if numbers.get("贴单费(CNY)") != 3:
            errors.append(f"定价依据第{row}行：贴单费必须为3元")
        required_math = {
            "采购价下限(CNY)", "采购价上限(CNY)",
            "国内运费下限(CNY)", "国内运费上限(CNY)",
            "贴单费(CNY)", "其他国内成本(CNY)",
            "成本价下限(CNY)", "成本价上限(CNY)",
        }
        if required_math.issubset(numbers):
            expected_low = (
                numbers["采购价下限(CNY)"] + numbers["国内运费下限(CNY)"]
                + numbers["贴单费(CNY)"] + numbers["其他国内成本(CNY)"]
            )
            expected_high = (
                numbers["采购价上限(CNY)"] + numbers["国内运费上限(CNY)"]
                + numbers["贴单费(CNY)"] + numbers["其他国内成本(CNY)"]
            )
            if abs(numbers["成本价下限(CNY)"] - expected_low) > 0.011:
                errors.append(f"定价依据第{row}行：成本价下限加总不正确")
            if abs(numbers["成本价上限(CNY)"] - expected_high) > 0.011:
                errors.append(f"定价依据第{row}行：成本价上限加总不正确")
            if numbers["成本价下限(CNY)"] > numbers["成本价上限(CNY)"]:
                errors.append(f"定价依据第{row}行：成本价下限高于上限")
        status = str(pricing.cell(row, columns["定价状态"]).value or "").strip()
        if not (
            status == "正式"
            or status.startswith("暂估：物流代理")
            or status.startswith("暂停")
        ):
            errors.append(
                f"定价依据第{row}行：定价状态必须为正式、暂估：物流代理或暂停"
            )
        international_freight = as_number(
            pricing.cell(row, columns["国际运费(CNY)"]).value
        )
        pricing_price = as_number(pricing.cell(row, columns["建议售价(RUB)"]).value)
        if international_freight is None and not status.startswith("暂停"):
            errors.append(f"定价依据第{row}行：非暂停状态必须有国际运费")
        elif international_freight is not None and international_freight < 0:
            errors.append(f"定价依据第{row}行：国际运费不能为负数")
        if pricing_price is None and not status.startswith("暂停"):
            errors.append(f"定价依据第{row}行：非暂停状态必须有建议售价")
        elif pricing_price is not None and pricing_price <= 0:
            errors.append(f"定价依据第{row}行：建议售价必须大于0")

        source = pricing.cell(row, columns["包装数据来源"]).value
        raw_or_derivation = pricing.cell(row, columns["包装数据原文/推导"]).value
        confidence = str(
            pricing.cell(row, columns["包装数据置信度"]).value or ""
        ).strip()
        recovery_note = pricing.cell(row, columns["缺失字段与补采记录"]).value
        if source in (None, ""):
            errors.append(f"定价依据第{row}行：包装数据来源为空")
        if raw_or_derivation in (None, ""):
            errors.append(f"定价依据第{row}行：包装数据原文/推导为空")
        if confidence not in {"高", "中", "低"}:
            errors.append(f"定价依据第{row}行：包装数据置信度必须为高/中/低")

        main_record = main_rows_by_item.get(item_key)
        if main_record:
            main_row, main_values = main_record
            main_price = as_number(main_values[8])
            row_complete = all(
                is_positive_number(main_values[index]) for index in (2, 3, 4, 5, 8)
            )
            if (main_price is None) != (pricing_price is None) or (
                main_price is not None
                and pricing_price is not None
                and abs(main_price - pricing_price) > 0.01
            ):
                errors.append(
                    f"自定义货号{item_key}：主表第{main_row}行建议售价"
                    f"与定价依据第{row}行不一致"
                )
            if row_complete and confidence == "低":
                errors.append(
                    f"定价依据第{row}行：完整SKU不能使用低置信度包装数据"
                )
            source_text = str(source or "")
            derivation_text = str(raw_or_derivation or "")
            uses_proxy = any(
                marker in source_text + derivation_text
                for marker in ("代理", "推导", "估算", "代理值")
            )
            if uses_proxy and not status.startswith("暂估：物流代理"):
                errors.append(
                    f"定价依据第{row}行：代理或推导包装数据必须标记暂估：物流代理"
                )
            if not row_complete and recovery_note in (None, ""):
                errors.append(
                    f"定价依据第{row}行：不完整SKU必须填写缺失字段与补采记录"
                )

    missing_pricing = sorted(set(main_rows_by_item) - set(pricing_rows_by_item))
    orphan_pricing = sorted(set(pricing_rows_by_item) - set(main_rows_by_item))
    if missing_pricing:
        errors.append("主表货号缺少定价依据行：" + "、".join(missing_pricing))
    if orphan_pricing:
        errors.append("定价依据存在主表没有的货号：" + "、".join(orphan_pricing))
    return errors


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    if len(sys.argv) != 2:
        print("用法: validate_listing_workbook.py <Ozon上品参考.xlsx>", file=sys.stderr)
        return 2
    path = Path(sys.argv[1]).expanduser().resolve()
    if not path.is_file():
        print(f"文件不存在: {path}", file=sys.stderr)
        return 2
    errors = validate(path)
    if errors:
        print("校验失败:")
        for error in errors:
            print(f"- {error}")
        return 1
    print("校验通过：九列位置、货币单位、详情链接和重量/建议售价完整率均符合要求。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
