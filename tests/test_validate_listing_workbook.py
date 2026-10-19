from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from validate_listing_workbook import (  # noqa: E402
    EXPECTED_HEADERS,
    REQUIRED_PRICING_HEADERS,
    validate,
)


def build_workbook(path: Path, complete_rows: int = 9, total_rows: int = 10) -> None:
    workbook = Workbook()
    main = workbook.active
    main.title = "Ozon上品参考"
    main.append(EXPECTED_HEADERS)
    pricing = workbook.create_sheet("定价依据")
    pricing.append(REQUIRED_PRICING_HEADERS)

    for index in range(total_rows):
        complete = index < complete_rows
        item = f"ITEM-{index + 1:02d}"
        price = 999 + index if complete else None
        main.append(
            [
                f"规格-{index + 1}",
                item,
                500 if complete else None,
                200 if complete else None,
                150 if complete else None,
                100 if complete else None,
                f"https://detail.1688.com/offer/{100000 + index}.html",
                20,
                price,
            ]
        )
        main.cell(main.max_row, 8).number_format = '¥0.00 "CNY"'
        if price is not None:
            main.cell(main.max_row, 9).number_format = '0 "RUB"'

        pricing.append(
            [
                item,
                10,
                12,
                2,
                3,
                3,
                1,
                16,
                19,
                25 if complete else None,
                price,
                "正式" if complete else "暂停：等待供应商包装数据",
                "1688详情页规格表",
                "包装尺寸200×150×100 mm，重量500 g" if complete else "页面未披露",
                "高" if complete else "低",
                "" if complete else "缺重量、长宽高、售价；已查规格区和详情图，待供应商确认",
            ]
        )
    workbook.save(path)


class ValidateListingWorkbookTests(unittest.TestCase):
    def make_path(self) -> Path:
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        return Path(temp_dir.name) / "listing.xlsx"

    def test_exactly_ninety_percent_complete_passes(self) -> None:
        path = self.make_path()
        build_workbook(path, complete_rows=9)
        self.assertEqual(validate(path), [])

    def test_eighty_percent_complete_fails(self) -> None:
        path = self.make_path()
        build_workbook(path, complete_rows=8)
        errors = validate(path)
        self.assertTrue(any("五项联合完整率不足90%" in error for error in errors))

    def test_duplicate_rows_are_rejected(self) -> None:
        path = self.make_path()
        build_workbook(path)
        workbook = load_workbook(path)
        main = workbook["Ozon上品参考"]
        main.cell(3, 2).value = main.cell(2, 2).value
        main.cell(3, 1).value = main.cell(2, 1).value
        main.cell(3, 7).value = main.cell(2, 7).value
        workbook.save(path)
        errors = validate(path)
        self.assertTrue(any("自定义货号" in error and "重复" in error for error in errors))
        self.assertTrue(any("详情链接与sku组合" in error for error in errors))

    def test_main_and_pricing_price_must_match_even_when_one_is_blank(self) -> None:
        path = self.make_path()
        build_workbook(path)
        workbook = load_workbook(path)
        pricing = workbook["定价依据"]
        pricing.cell(2, REQUIRED_PRICING_HEADERS.index("建议售价(RUB)") + 1).value = None
        workbook.save(path)
        errors = validate(path)
        self.assertTrue(any("建议售价" in error and "不一致" in error for error in errors))

    def test_negative_and_proxy_data_without_provisional_status_fail(self) -> None:
        path = self.make_path()
        build_workbook(path)
        workbook = load_workbook(path)
        pricing = workbook["定价依据"]
        pricing.cell(2, REQUIRED_PRICING_HEADERS.index("国内运费下限(CNY)") + 1).value = -1
        pricing.cell(2, REQUIRED_PRICING_HEADERS.index("包装数据来源") + 1).value = "同供应商代理值"
        workbook.save(path)
        errors = validate(path)
        self.assertTrue(any("不能为负数" in error for error in errors))
        self.assertTrue(any("必须标记暂估：物流代理" in error for error in errors))

    def test_missing_and_orphan_pricing_rows_fail(self) -> None:
        path = self.make_path()
        build_workbook(path)
        workbook = load_workbook(path)
        pricing = workbook["定价依据"]
        pricing.cell(2, 1).value = "ORPHAN"
        workbook.save(path)
        errors = validate(path)
        self.assertTrue(any("缺少定价依据行" in error for error in errors))
        self.assertTrue(any("主表没有的货号" in error for error in errors))


if __name__ == "__main__":
    unittest.main()
